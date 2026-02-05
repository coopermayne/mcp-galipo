"""
Split the master pleading template into a base + fragment files.

Takes the styled master template and creates:
- base_subdoc.docx: Main template with subdoc placeholders
- fragments/notice.docx: Notice of Motion + Meet & Confer
- fragments/toc.docx: Table of Contents
- fragments/toa.docx: Table of Authorities
- fragments/cert_compliance.docx: Certificate of Compliance

All fragments preserve the original styles since they're derived
from the same master template.

Usage:
    python scripts/split_template.py
"""

import copy
from pathlib import Path
from lxml import etree
from docx import Document
from docx.oxml.ns import qn

MASTER = Path(__file__).parent.parent / "templates" / "pleading_template_202601151640.docx"
OUTPUT_DIR = Path(__file__).parent.parent / "templates" / "pleadings"
FRAGMENTS_DIR = OUTPUT_DIR / "fragments"


def get_element_text(el):
    """Get all text from an element (for debugging)."""
    texts = []
    for t in el.iter(qn('w:t')):
        if t.text:
            texts.append(t.text)
    return ''.join(texts).strip()


def delete_elements(body, indices):
    """Delete elements from body by index (handles reverse order)."""
    elements = list(body)
    for idx in sorted(indices, reverse=True):
        body.remove(elements[idx])


def create_fragment(master_path, keep_indices, output_path):
    """
    Create a fragment by copying master and keeping only specified elements.

    This preserves all style definitions from the master.
    """
    doc = Document(str(master_path))
    body = doc.element.body
    elements = list(body)

    # Remove everything except the kept indices and sectPr
    for idx in range(len(elements) - 1, -1, -1):
        tag = elements[idx].tag.split('}')[-1]
        if tag == 'sectPr':
            continue  # Always keep section properties
        if idx not in keep_indices:
            body.remove(elements[idx])

    doc.save(str(output_path))
    print(f"  Created: {output_path.name} ({len(keep_indices)} elements)")


def create_base(master_path, output_path):
    """
    Create the base template by removing fragment sections
    and inserting subdoc placeholders.
    """
    doc = Document(str(master_path))
    body = doc.element.body
    elements = list(body)

    # Elements to remove (fragment content):
    # Notice: 4-16 (includes {%p if/endif %})
    # TOC: 17-22 (page break + heading + entries)
    # TOA: 23-25 (page break + heading + entries)
    # Cert: 44-55 (includes {%p if/endif %})
    # Bookmark ends: 37-42
    remove_indices = set(range(4, 26)) | set(range(37, 43)) | set(range(44, 56))

    # First, identify where to insert placeholders (before removing)
    # We'll insert after element 3 (caption table) for notice/toc/toa
    # and after element 43 (attorneys line) for cert

    # Remove elements in reverse order
    for idx in sorted(remove_indices, reverse=True):
        body.remove(elements[idx])

    # Now the body has:
    # 0: Attorney table
    # 1: Court name
    # 2: Empty para
    # 3: Caption table
    # (removed 4-25: notice, toc, toa)
    # 4 (was 26): MEMORANDUM heading
    # ...
    # (removed cert)

    # Insert subdoc placeholder paragraphs after the caption table (element 3)
    # We need to create proper paragraph elements with Jinja2 tags
    caption_table = list(body)[3]

    # Create placeholder paragraphs (insert in reverse order since each goes after caption)
    # {%p if %} and {%p endif %} must be on SEPARATE paragraphs for docxtpl
    placeholder_groups = [
        ("include_toa", "subdoc_toa"),
        ("include_toc", "subdoc_toc"),
        ("include_meet_confer", "subdoc_meet_confer"),
        ("include_notice", "subdoc_notice"),
    ]

    for condition, subdoc_var in placeholder_groups:
        endif_el = make_paragraph("{%p endif %}", "normal0")
        subdoc_el = make_paragraph("{{ " + subdoc_var + " }}", "normal0")
        if_el = make_paragraph("{%p if " + condition + " %}", "normal0")
        # Insert in reverse order (endif, subdoc, if) since addnext puts after
        caption_table.addnext(endif_el)
        caption_table.addnext(subdoc_el)
        caption_table.addnext(if_el)

    # Insert cert_compliance placeholder after the main signature block
    # Find the "Attorneys for" paragraph (should be after conclusion/sig block)
    current_elements = list(body)
    # Find the last "Attorneys for" paragraph in the memo signature area
    attorneys_idx = None
    for idx, el in enumerate(current_elements):
        tag = el.tag.split('}')[-1]
        if tag == 'p':
            text = get_element_text(el)
            if 'Attorneys for' in text:
                attorneys_idx = idx

    if attorneys_idx is not None:
        # {%p if %} and {%p endif %} must be on SEPARATE paragraphs
        cert_endif = make_paragraph("{%p endif %}", "normal0")
        cert_subdoc = make_paragraph("{{ subdoc_cert_compliance }}", "normal0")
        cert_if = make_paragraph("{%p if include_cert_compliance %}", "normal0")
        current_elements[attorneys_idx].addnext(cert_endif)
        current_elements[attorneys_idx].addnext(cert_subdoc)
        current_elements[attorneys_idx].addnext(cert_if)

    doc.save(str(output_path))
    print(f"  Created: {output_path.name}")


def make_paragraph(text, style_name=None):
    """Create a paragraph XML element with text and optional style."""
    nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    p = etree.SubElement(etree.Element('dummy'), qn('w:p'))

    if style_name:
        pPr = etree.SubElement(p, qn('w:pPr'))
        pStyle = etree.SubElement(pPr, qn('w:pStyle'))
        pStyle.set(qn('w:val'), style_name)

    r = etree.SubElement(p, qn('w:r'))
    t = etree.SubElement(r, qn('w:t'))
    t.text = text
    t.set(qn('xml:space'), 'preserve')

    return p


def main():
    FRAGMENTS_DIR.mkdir(exist_ok=True)

    print(f"Master template: {MASTER.name}")
    print()

    # Verify the master exists
    if not MASTER.exists():
        print(f"ERROR: Master template not found: {MASTER}")
        return

    # Create fragments
    print("Creating fragments...")

    # Notice: elements 5-15 (skip the {%p if %} at 4 and {%p endif %} at 16)
    create_fragment(MASTER, set(range(5, 16)), FRAGMENTS_DIR / "notice.docx")

    # TOC: elements 17-22 (page break + heading + entries)
    create_fragment(MASTER, set(range(17, 23)), FRAGMENTS_DIR / "toc.docx")

    # TOA: elements 23-25 (page break heading + heading + entries)
    create_fragment(MASTER, set(range(23, 26)), FRAGMENTS_DIR / "toa.docx")

    # Meet & Confer is part of the notice section in the original.
    # For now, notice.docx includes both the notice and meet & confer text.
    # We can split them later if needed.
    # Create a standalone meet_confer fragment from element 8 only
    create_fragment(MASTER, {8}, FRAGMENTS_DIR / "meet_confer.docx")

    # Cert compliance: elements 45-54 (skip the {%p if %} at 44 and {%p endif %} at 55)
    create_fragment(MASTER, set(range(45, 55)), FRAGMENTS_DIR / "cert_compliance.docx")

    print()

    # Create base template
    print("Creating base template...")
    create_base(MASTER, OUTPUT_DIR / "base_subdoc.docx")

    print()
    print("Done! Template split complete.")
    print()
    print("Files created:")
    print(f"  {OUTPUT_DIR / 'base_subdoc.docx'}")
    for f in sorted(FRAGMENTS_DIR.glob("*.docx")):
        print(f"  {f}")


if __name__ == "__main__":
    main()
